from pathlib import Path
import sys,hashlib,json,re,zipfile
sys.path.insert(0,str(Path('.local/presentation-tools').resolve()))
from pptx import Presentation
from PIL import Image,ImageDraw
from lxml import etree
out=Path('presentation/v2');src=Path('presentation-deck.pptx');r=Presentation('presentation-deck-v2.pptx');source=Presentation(out/'user-edited-source.pptx')
assert hashlib.sha256(src.read_bytes()).hexdigest()==(out/'source-sha256.txt').read_text().strip()
assert len(r.slides)==13
assert all(str(s.background.fill.fore_color.rgb)=='080F1B' for s in r.slides)
assert len([q.text for q in r.slides[0].shapes if q.has_text_frame and q.text.strip()])==1
assert not any(q.shape_type==13 for q in r.slides[1].shapes)
assert not any(re.fullmatch(r'\d+\s*/\s*\d+',q.text.strip()) for s in r.slides for q in s.shapes if q.has_text_frame)
preservation=[]
for oldidx,newidx in [(3,3),(4,4),(5,5),(6,6),(7,7),(8,8),(9,10)]:
 old=[]
 for q in source.slides[oldidx].shapes:
  if not q.has_text_frame or not q.text.strip():continue
  v=q.text.strip();y=q.top/914400
  if y<.65 or re.fullmatch(r'\d+\s*/\s*\d+',v) or (v=='FINBRIDGE AI' and y>6.8):continue
  old.append(v)
 current={q.text.strip() for q in r.slides[newidx].shapes if q.has_text_frame}
 missing=[v for v in old if v not in current];assert not missing,(oldidx,missing)
 preservation.append({'source_slide':oldidx+1,'v2_slide':newidx+1,'retained_text_elements':len(old)})
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with zipfile.ZipFile('presentation-deck-v2.pptx') as z:
 for name in z.namelist():
  if re.fullmatch(r'ppt/slides/slide\d+\.xml',name):
   doc=etree.fromstring(z.read(name));assert doc.find('p:timing',ns) is not None
   ids=set(doc.xpath('//p:cNvPr/@id',namespaces=ns));assert all(x in ids for x in doc.xpath('//p:spTgt/@spid',namespaces=ns))
layout=json.loads((out/'layout-report.json').read_text(encoding='utf-8-sig'));assert layout==[]
refs=[]
for path in Path('src').rglob('*'):
 if path.is_file() and 'assets' not in path.parts and path.suffix in ['.ts','.tsx','.css','.json']:
  if 'finbridge-logo' in path.read_text(encoding='utf-8',errors='replace'):refs.append(str(path))
assert not refs
logo=Image.open('src/assets/finbridge-logo.png');assert logo.mode=='RGBA' and logo.getpixel((0,0))[3]==0
result={'slide_count':13,'source_unchanged':True,'developer_position':2,'developer_avatars':0,'cover_has_only_product_title':True,'slide_numbers_removed':True,'all_backgrounds_dark':True,'logo_application_references':refs,'powerpoint_text_overflow':layout,'preserved_source_content':preservation}
(out/'validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
files=sorted((out/'preview').glob('*.PNG'),key=lambda p:int(''.join(c for c in p.stem if c.isdigit())))
im=Image.new('RGB',(1600,474*((len(files)+1)//2)),'#172A40')
for i,p in enumerate(files):
 t=Image.open(p).convert('RGB');t.thumbnail((780,439));x=10+i%2*800;y=10+i//2*474;im.paste(t,(x,y));ImageDraw.Draw(im).text((x+8,y+444),f'Slide {i+1}',fill='#EDF5FF')
im.save(out/'preview/contact-sheet.jpg',quality=94)
print(json.dumps(result))
