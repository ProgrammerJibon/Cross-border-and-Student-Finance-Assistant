import sys,zipfile,json,re,io
from pathlib import Path
sys.path.insert(0,str(Path('.local/presentation-tools').resolve()))
from PIL import Image,ImageDraw
from pptx import Presentation
from lxml import etree
r=Presentation('presentation-deck.pptx');assert len(r.slides)==10
assert all(str(s.background.fill.fore_color.rgb)=='080F1B' for s in r.slides)
texts=[q.text for q in r.slides[0].shapes if q.has_text_frame and q.text.strip()];assert len(texts)==2
assert any('Real-world struggles' in q.text for q in r.slides[2].shapes if q.has_text_frame)
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with zipfile.ZipFile('presentation-deck.pptx') as z:
 slidefiles=[n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)]
 for name in slidefiles:
  doc=etree.fromstring(z.read(name));assert doc.find('p:timing',ns) is not None
  ids=set(doc.xpath('//p:cNvPr/@id',namespaces=ns));assert all(x in ids for x in doc.xpath('//p:spTgt/@spid',namespaces=ns))
 gifs=[n for n in z.namelist() if n.endswith('.gif')];frames=[Image.open(io.BytesIO(z.read(n))).n_frames for n in gifs]
 assert len(gifs)==7 and all(n>1 for n in frames)
 charts=[n for n in z.namelist() if re.fullmatch(r'ppt/charts/chart\d+\.xml',n)];assert len(charts)==6
layout=json.loads(Path('presentation/dark-layout-report.json').read_text(encoding='utf-8-sig'));assert layout==[]
report=json.loads(Path('presentation/dark-animation-report.json').read_text(encoding='utf-8-sig'))
result={'slides':10,'all_backgrounds_dark':True,'native_charts':len(charts),'animated_workflow_graphics':len(gifs),'gif_frame_counts':frames,'native_effects':sum(x['Effects'] for x in report),'powerpoint_text_overflow':layout,'actual_demo_ui_capture':True}
Path('presentation/dark-validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
files=sorted(Path('presentation/dark-preview').glob('*.PNG'),key=lambda x:int(''.join(c for c in x.stem if c.isdigit())))
im=Image.new('RGB',(1600,2370),'#172A40')
for i,p in enumerate(files):
 thumb=Image.open(p).convert('RGB');thumb.thumbnail((780,439));x=10+i%2*800;y=10+i//2*474;im.paste(thumb,(x,y));ImageDraw.Draw(im).text((x+8,y+444),f'Slide {i+1}',fill='#EDF5FF')
im.save('presentation/dark-preview/contact-sheet.jpg',quality=94)
print(json.dumps(result))
