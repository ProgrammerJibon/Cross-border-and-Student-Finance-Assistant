from pathlib import Path
import sys, math, json
sys.path.insert(0,str(Path('.local/presentation-tools').resolve()))
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.xmlchemy import OxmlElement

BASE=Path('presentation');AS=BASE/'visuals';AS.mkdir(exist_ok=True)
old=Presentation(str(BASE/'previous-deck.pptx'))
OLD_NOTES=[s.notes_slide.notes_text_frame.text for s in old.slides]
P=Presentation();P.slide_width=Inches(13.333);P.slide_height=Inches(7.5)
P.core_properties.title='FinBridge AI';P.core_properties.author='FinBridge AI Project Team';P.core_properties.subject='Visual project presentation'
INK='102F36';BG='F5F6F1';TEAL='008D78';MINT='57DBBA';BLUE='749FF5';GOLD='F1BA61';GREY='617579';PALE='E6ECE6';WHITE='FFFFFF';RED='DF7A71';DARK='102F36'

def rgb(c):return RGBColor.from_string(c)
def shape(s,x,y,w,h,fill=PALE,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
 q=s.shapes.add_shape(kind, Inches(x),Inches(y), Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=rgb(fill);q.line.fill.background()
 if line:q.line.color.rgb=rgb(line)
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:q.adjustments[0]=.12
 return q

def text(s,t,x,y,w,h,size=18,c=INK,b=False,align=None,reveal=None):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,ln in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=ln;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=b;p.font.color.rgb=rgb(c);p.space_after=Pt(4)
  if align is not None:p.alignment=align
 if reveal is not None:q.name=f'reveal_{reveal:02d}'
 return q

def circle(s,x,y,d,c):return shape(s,x,y,d,d,c,MSO_SHAPE.OVAL)
def ln(s,x1,y1,x2,y2,c=PALE,width=1):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2));q.line.color.rgb=rgb(c);q.line.width=Pt(width);return q

def new(title=None,section=None):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 n=len(P.slides)
 if title:
  text(s,section.upper(),.65,.42,11,.25,10,TEAL,True)
  text(s,title,.65,.94,12,1.0,34,INK,True,reveal=1)
  text(s,'FinBridge AI',.65,7.05,3,.23,10,GREY)
  text(s,f'{n:02d}',12.02,6.99,.65,.3,12,TEAL,True,PP_ALIGN.RIGHT)
 return s

def note(s,n,extra=''):
 s.notes_slide.notes_text_frame.text=OLD_NOTES[n-1].split('\nAnimation:')[0]+'\n\nRedesigned visual edition. '+extra+' Illustrative charts are not live financial data or measured product outcomes. Animations: distinct looping conceptual motion graphic plus native staggered entrance effects, with fade transition. Feature labels are concise; this note retains detailed feature coverage.'

# Hand-drawn line icons, repeated consistently across the visual system.
def icon(d,kind,x,y,sc=1,c='#57DBBA',lw=5):
 def box(a,b,cc=None,r=8):d.rounded_rectangle((x+a[0]*sc,y+a[1]*sc,x+b[0]*sc,y+b[1]*sc),radius=r*sc,outline=cc or c,width=max(2,int(lw*sc)))
 def line(pts,cc=None):d.line([(x+a*sc,y+b*sc) for a,b in pts],fill=cc or c,width=max(2,int(lw*sc)),joint='curve')
 def ell(a,b):d.ellipse((x+a[0]*sc,y+a[1]*sc,x+b[0]*sc,y+b[1]*sc),outline=c,width=max(2,int(lw*sc)))
 if kind=='wallet':
  box((8,20),(86,76));box((57,36),(92,59));ell((67,44),(73,50));line([(18,19),(18,10),(72,10),(72,20)])
 elif kind=='student':
  line([(3,28),(48,8),(93,28),(48,48),(3,28)]);line([(20,38),(20,60),(48,72),(76,60),(76,38)]);line([(92,28),(92,65)])
 elif kind=='family':
  ell((10,10),(38,38));ell((56,10),(84,38));d.arc((x,y+38*sc,x+48*sc,y+87*sc),180,360,fill=c,width=max(2,int(lw*sc)));d.arc((x+45*sc,y+38*sc,x+94*sc,y+87*sc),180,360,fill=c,width=max(2,int(lw*sc)))
 elif kind=='shield':
  line([(48,4),(86,20),(82,56),(69,76),(48,91),(27,76),(14,56),(10,20),(48,4)]);line([(28,45),(42,60),(68,32)])
 elif kind=='ai':
  box((17,17),(77,77),r=13)
  for v in [30,48,65]:
   line([(v,5),(v,17)]);line([(v,77),(v,89)]);line([(5,v),(17,v)]);line([(77,v),(89,v)])
  ell((31,34),(40,43));ell((55,34),(64,43));line([(32,60),(63,60)])
 elif kind=='chart':
  line([(8,9),(8,82),(92,82)])
  for a,b in [(22,49),(44,33),(66,14)]:box((a,b),(a+13,74),r=2)
 elif kind=='tuition':
  line([(6,28),(48,7),(90,28),(6,28)]);line([(9,82),(88,82)])
  for v in [21,45,70]:line([(v,39),(v,69)])
 elif kind=='clock':
  ell((9,9),(87,87));line([(48,25),(48,49),(66,59)])
 elif kind=='transfer':
  line([(7,28),(83,28),(65,11)]);line([(83,28),(65,45)]);line([(88,68),(12,68),(30,51)]);line([(12,68),(30,85)])
 elif kind=='lock':
  box((16,40),(80,88));d.arc((x+29*sc,y+5*sc,x+67*sc,y+62*sc),180,360,fill=c,width=max(2,int(lw*sc)));ell((43,56),(53,66));line([(48,65),(48,74)])
 elif kind=='database':
  ell((12,8),(84,34));d.arc((x+12*sc,y+58*sc,x+84*sc,y+84*sc),0,180,fill=c,width=max(2,int(lw*sc)));line([(12,21),(12,71)]);line([(84,21),(84,71)]);d.arc((x+12*sc,y+33*sc,x+84*sc,y+59*sc),0,180,fill=c,width=max(2,int(lw*sc)))
 elif kind=='check':line([(13,48),(38,71),(85,20)])
 elif kind=='document':
  box((18,6),(78,89),r=5)
  for v in [29,46,63]:line([(30,v),(65,v)])
 elif kind=='globe':
  ell((5,5),(90,90));ell((27,5),(68,90));line([(5,47),(90,47)])
 else:icon(d,'chart',x,y,sc,c,lw)

def card(d,box,fill='#1A4147',outline=None):d.rounded_rectangle(box,radius=22,fill=fill,outline=outline,width=2)
def motion(kind,n,w=960,h=720):
 path=AS/f'{n:02d}-{kind}.gif';frames=[]
 for k in range(32):
  t=k/32;im=Image.new('RGB',(w,h),'#'+DARK);d=ImageDraw.Draw(im)
  # quiet coordinate grid, part of the visual language
  for xx in range(40,w,70):
   for yy in range(40,h,70):d.ellipse((xx,yy,xx+2,yy+2),fill='#285058')
  def particle(a,b,phase=0,col='#57DBBA'):
   u=(t+phase)%1;x=a[0]+(b[0]-a[0])*u;y=a[1]+(b[1]-a[1])*u;d.ellipse((x-7,y-7,x+7,y+7),fill=col)
  if kind in ['bridge','team']:
   # Arc joining two financial worlds; growing bars beneath.
   d.arc((120,100,w-120,h+100),190,350,fill='#57DBBA',width=7)
   d.line((110,h*.66,w-110,h*.66),fill='#42706F',width=5)
   for j in range(7):
    x=180+j*(w-360)/6;hh=65+j*17+12*math.sin(t*math.tau+j);d.rounded_rectangle((x,h*.81-hh,x+36,h*.81),radius=8,fill='#315E61')
   for x in [110,w-220]:
    card(d,(x,h*.47,x+110,h*.47+112));icon(d,'wallet' if x==110 else 'student',x+9,h*.47+9,1)
   u=t;px=180+(w-360)*u;py=h*.40-120*math.sin(u*math.pi)
   d.ellipse((px-23,py-23,px+23,py+23),fill='#F1BA61');d.line((px,py-12,px,py+12),fill='#102F36',width=4)
   if kind=='bridge':
    card(d,(w*.40,h*.56,w*.65,h*.87));icon(d,'ai',w*.47,h*.60,1.0,'#749FF5')
  elif kind=='problem':
   centers=[(230,200),(690,200),(230,520),(690,520)]
   for i,((cx,cy),ic) in enumerate(zip(centers,['wallet','clock','transfer','family'])):
    yy=cy+int(10*math.sin(t*math.tau+i));card(d,(cx-115,yy-95,cx+115,yy+95));icon(d,ic,cx-48,yy-48,1.05,'#F1BA61' if i==1 else '#57DBBA')
    for j in range(5):
     xx=cx+(480-cx)*(j/6);py=yy+(360-yy)*(j/6);d.ellipse((xx,py,xx+4,py+4),fill='#42616A')
   d.ellipse((414,294,546,426),outline='#DF7A71',width=4)
   d.line((480,325,480,369),fill='#DF7A71',width=9);d.ellipse((475,389,485,399),fill='#DF7A71')
  elif kind=='solution':
   center=(480,360);points=[(480,140),(230,510),(730,510)]
   for i,(p,ic) in enumerate(zip(points,['student','family','shield'])):
    d.line((*center,*p),fill='#3F696C',width=4);particle(p,center,i*.28)
    d.ellipse((p[0]-80,p[1]-80,p[0]+80,p[1]+80),fill='#1A4147');icon(d,ic,p[0]-48,p[1]-48,1)
   rr=104+9*math.sin(t*math.tau);d.ellipse((480-rr,360-rr,480+rr,360+rr),outline='#315F61',width=2)
   d.ellipse((390,270,570,450),fill='#57DBBA');icon(d,'ai',432,312,1,'#102F36')
  elif kind=='ai':
   for i,ic in enumerate(['wallet','document','chart']):
    y=110+i*210;card(d,(80,y,235,y+140));icon(d,ic,110,y+22,1)
    d.line((235,y+70,480,360),fill='#3D686B',width=4);particle((235,y+70),(480,360),i/3)
   d.ellipse((376,256,584,464),fill='#57DBBA');icon(d,'ai',432,312,1,'#102F36')
   d.line((584,360,752,360),fill='#3D686B',width=4);particle((584,360),(752,360),.2,'#F1BA61')
   card(d,(730,205,900,515));icon(d,'document',766,237,1,'#F1BA61');icon(d,'check',775,393,.8,'#57DBBA')
  elif kind=='pipeline':
   yy=h*.50;xs=[120,350,580,810,1040]
   for j in range(4):d.line((xs[j],yy,xs[j+1],yy),fill='#426C6D',width=5);particle((xs[j],yy),(xs[j+1],yy),-j*.17)
   for i,(xx,ic) in enumerate(zip(xs,['family','ai','check','shield','database'])):
    d.ellipse((xx-72,yy-72,xx+72,yy+72),fill='#1C474D' if i!=2 else '#57DBBA');icon(d,ic,xx-46,yy-46,.96,'#102F36' if i==2 else '#57DBBA')
  elif kind=='student':
   card(d,(80,85,880,635));card(d,(105,110,855,193),'#234C52');icon(d,'wallet',132,122,.6)
   for i in range(3):d.rounded_rectangle((530+i*90,137,590+i*90,154),radius=7,fill='#315F65')
   # donut + chart inside an abstract dashboard
   d.arc((145,260,405,520),-90,270,fill='#315E64',width=32);d.arc((145,260,405,520),-90,140+8*math.sin(t*math.tau),fill='#57DBBA',width=32);icon(d,'student',225,338,1)
   pts=[]
   for j in range(8):
    xx=475+j*43;yy=470-j*20+math.sin(j*1.4+t*math.tau)*14;pts.append((xx,yy));d.line((xx,518,xx,yy+25),fill='#2D555D',width=18)
   d.line(pts,fill='#F1BA61',width=6);cx,cy=pts[k%8];d.ellipse((cx-8,cy-8,cx+8,cy+8),fill='#FFFFFF')
   for i in range(3):d.rounded_rectangle((160+i*220,566,320+i*220,579),radius=6,fill=['#57DBBA','#749FF5','#F1BA61'][i])
  elif kind=='parent':
   d.ellipse((105,240,285,420),fill='#224B50');icon(d,'family',147,282,1)
   d.line((285,330,410,330),fill='#426B6F',width=5);particle((285,330),(410,330))
   card(d,(400,105,850,595));icon(d,'shield',574,137,1.1)
   for i,(val,col) in enumerate([(0.83,'#749FF5'),(.66,'#57DBBA'),(.42,'#F1BA61')]):
    yy=335+i*81;d.rounded_rectangle((455,yy,795,yy+12),radius=6,fill='#355E64');xx=455+340*(val+.015*math.sin(t*math.tau));d.rounded_rectangle((455,yy,xx,yy+12),radius=6,fill=col);d.ellipse((xx-15,yy-9,xx+15,yy+21),fill=col)
  elif kind=='admin':
   for i,ic in enumerate(['family','ai','transfer']):
    xx=90+i*270;card(d,(xx,90,xx+240,270));icon(d,ic,xx+75,125,.95)
   card(d,(90,310,600,620));card(d,(630,310,870,620))
   for j in range(6):
    xx=135+j*70;hh=80+j*24+8*math.sin(t*math.tau+j);d.rounded_rectangle((xx,575-hh,xx+38,575),radius=7,fill='#57DBBA' if j%2 else '#749FF5')
   for j in range(4):
    yy=352+j*65;d.ellipse((660,yy,679,yy+19),fill='#57DBBA' if (k//8)%4==j else '#42666C');d.rounded_rectangle((702,yy+4,829,yy+13),radius=4,fill='#42666C')
  elif kind=='security':
   center=(480,360);pts=[(230,150),(730,150),(230,570),(730,570)]
   for i,(p,ic) in enumerate(zip(pts,['lock','ai','family','database'])):
    d.line((*p,*center),fill='#3F666C',width=3);particle(p,center,i/4);d.ellipse((p[0]-66,p[1]-66,p[0]+66,p[1]+66),fill='#1A4147');icon(d,ic,p[0]-37,p[1]-37,.78)
   r=154+8*math.sin(t*math.tau);d.ellipse((480-r,360-r,480+r,360+r),outline='#2B575B',width=3)
   icon(d,'shield',380,257,2.1,'#57DBBA',4)
  frames.append(im.quantize(colors=80,method=Image.Quantize.MEDIANCUT))
 frames[0].save(path,save_all=True,append_images=frames[1:],duration=65,loop=0,optimize=True,disposal=2)
 return path

def picture(s,path,x,y,w,h,order=2):
 p=s.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h));p.name=f'reveal_{order:02d}';return p

def art(s,kind,n,x,y,w,h,order=2):
 return picture(s,motion(kind,n,1160 if kind=='pipeline' else 960,420 if kind=='pipeline' else 720),x,y,w,h,order)

def native_chart(s,kind,x,y,w,h,cats,vals,colors,series='Illustrative data',order=4):
 data=CategoryChartData();data.categories=cats;data.add_series(series,vals)
 typ={'donut':XL_CHART_TYPE.DOUGHNUT,'bar':XL_CHART_TYPE.BAR_CLUSTERED,'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'line':XL_CHART_TYPE.LINE}[kind]
 sh=s.shapes.add_chart(typ, Inches(x), Inches(y), Inches(w), Inches(h), data);sh.name=f'reveal_{order:02d}';ch=sh.chart;ch.has_legend=False;ch.has_title=False;ch.font.name='Aptos';ch.font.size=Pt(11);ch.font.color.rgb=rgb(GREY)
 ser=ch.series[0]
 for i,pt in enumerate(ser.points):pt.format.fill.solid();pt.format.fill.fore_color.rgb=rgb(colors[i%len(colors)]);pt.format.line.fill.background()
 if kind=='donut':ch.plots[0].hole_size=75;ch.plots[0].first_slice_angle=270
 else:
  ch.category_axis.tick_labels.font.size=Pt(10);ch.category_axis.tick_labels.font.color.rgb=rgb(GREY);ch.value_axis.tick_labels.font.size=Pt(9);ch.value_axis.tick_labels.font.color.rgb=rgb(GREY);ch.value_axis.has_major_gridlines=False
  ch.category_axis.format.line.fill.background();ch.value_axis.format.line.fill.background()
  if kind=='bar':ch.value_axis.visible=False;ch.plots[0].gap_width=60;ch.plots[0].has_data_labels=True;ch.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;ch.plots[0].data_labels.font.size=Pt(10);ch.plots[0].data_labels.font.color.rgb=rgb(INK)
 return sh

def point(s,num,label,detail,x,y,w=5.0,order=3):
 circle(s,x,y+.03,.35,TEAL);text(s,str(num),x+.01,y+.075,.33,.22,10,WHITE,True,PP_ALIGN.CENTER)
 text(s,label,x+.56,y,w-.56,.40,23,INK,True,reveal=order)
 if detail:text(s,detail,x+.56,y+.49,w-.56,.60,15,GREY,reveal=order)

def feature(s,ic,title,body,x,y,w=3.35,order=3):
 # Small icon is a reusable transparent PNG, not an invented portrait.
 ip=AS/f'icon-{ic}.png'
 if not ip.exists():
  im=Image.new('RGBA',(112,112),(0,0,0,0));icon(ImageDraw.Draw(im),ic,7,7,1,'#'+TEAL,4);im.save(ip)
 picture(s,ip,x,y,.35,.35,order)
 text(s,title,x+.48,y+.01,w-.48,.35,16,INK,True,reveal=order)
 text(s,body,x+.48,y+.49,w-.48,1.00,12.5,GREY,reveal=order)

# 1. Only two textual elements; the artwork contains no labels.
s=new();s.background.fill.fore_color.rgb=rgb(DARK)
text(s,'FinBridge AI',.70,2.05,6.1,.95,50,WHITE,True,reveal=1)
text(s,'An AI copilot for student money,\nfamily funding and cross-border finance.',.74,3.20,5.8,1.35,24,'B4D0CF',reveal=2)
art(s,'bridge',1,6.55,1.05,6.4,4.8,3)
note(s,1,'The title slide has exactly the product name and a single-sentence description. Large animated bridge visual represents family funding reaching study goals.')

# 2. Four compact problems; most canvas is the visual.
s=new('Money is connected. The tools aren’t.','Problem statement')
art(s,'problem',2,6.52,2.00,6.16,4.62)
for i,(a,b) in enumerate([('Scattered balances','Accounts and currencies'),('Tuition uncertainty','Deadlines and living costs'),('Hidden transfer costs','Fees and exchange spreads'),('Family coordination','Visibility and approvals')]):point(s,i+1,a,b,.72,2.17+i*1.10,5.45,3+i)
note(s,2,'Four disconnected visual nodes represent the actual product design problems, not market survey statistics.')

# 3
s=new('One financial picture. Three connected roles.','Our solution')
art(s,'solution',3,6.52,2.00,6.16,4.62)
for i,(a,b) in enumerate([('Students plan','Money, spending and tuition'),('Parents support','Funding, visibility and limits'),('Admins oversee','Policies, settlements and evidence')]):point(s,i+1,a,b,.72,2.27+i*1.31,5.55,3+i)
text(s,'AI connects insight to reviewed action.',1.28,6.35,5.15,.40,17,TEAL,True,reveal=6)
note(s,3,'Shared capabilities: English, Bengali and Chinese; light/dark/system theme; global search; profile, locale, country, currency and time-zone preferences. The orbit diagram places AI at the intersection of the three role experiences.')

# 4
s=new('Ask. Understand. Act.','Our implementation / AI copilot')
art(s,'ai',4,.65,2.03,6.0,4.5)
text(s,'“Can I afford next semester?”',7.03,2.08,5.6,.65,24,INK,True,reveal=3)
text(s,'Reads your data → calculates the gap\n→ prepares a plan for your review.',7.05,2.95,5.45,.83,19,GREY,reveal=4)
text(s,'Illustrative tuition plan · CNY',7.05,4.05,5.45,.28,11,TEAL,True)
native_chart(s,'bar',6.96,4.42,5.65,1.52,['Tuition','Available','Gap'],[42000,30000,12000],[INK,MINT,GOLD],order=5)
text(s,'Gemini  /  OpenAI  /  local mock',7.05,6.31,5.6,.35,15,TEAL,True,reveal=6)
note(s,4,'The illustrative chart shows 42,000 CNY tuition, 30,000 CNY available funding and a 12,000 CNY gap. Capabilities also include budget recommendations, FX/provider comparison, financial health, credit guidance, follow-up questions, navigation, payment drafts and role-checked action proposals. The model can prepare drafts and other allowed low-impact tool outputs; financial execution requires confirmation.')

# 5
s=new('AI proposes. People approve. The server verifies.','Our implementation / Execution flow')
art(s,'pipeline',5,.65,2.05,12.03,4.35)
# Labels sit beneath icons inside the dark diagram, all editable.
for i,(a,b) in enumerate([('Request','Student / parent'),('AI tools','Structured schemas'),('Review','User confirmation'),('Server','Auth + risk + policy'),('Ledger','Atomic commit')]):
 x=.65+(120+i*230)/1160*12.03
 text(s,a,x-1.0,5.27,2,.35,20,WHITE,True,PP_ALIGN.CENTER,reveal=3+i)
 text(s,b,x-1.05,5.78,2.1,.30,11,'B4D0CF',False,PP_ALIGN.CENTER,reveal=3+i)
text(s,'Next.js + React',.72,6.57,3,.32,14,TEAL,True)
text(s,'TypeScript + Zod',4.0,6.57,3,.32,14,TEAL,True)
text(s,'MySQL + Decimal.js',8.0,6.57,4.6,.32,14,TEAL,True,PP_ALIGN.RIGHT)
note(s,5,'The five stages animate from user intent to persisted outcome. Shared payment lifecycle: draft, parent review when required, queued, processing, completed/failed/cancelled. External channels use manual settlement. Confirmations reuse application Server Actions and shared payment services. Database has 57 tables; row locks and idempotency guard against duplicate execution. All monetary activity is sandbox.')

# 6 student
s=new('Student features','Features / Student')
art(s,'student',6,.65,2.04,4.22,3.165)
text(s,'Plan every semester.',.72,5.47,4.15,.55,24,INK,True,reveal=3)
native_chart(s,'donut',.66,6.06,1.15,.75,['Housing','Food','Transport','Savings'],[45,18,17,20],[INK,TEAL,BLUE,GOLD],order=4)
text(s,'Budget mix\nIllustrative',1.92,6.12,2.8,.56,12,GREY)
student=[
 ('wallet','Money & transactions','Multi-currency wallets; cash; income/expenses; categories; search/filter; CSV mapping and preview.'),
 ('tuition','Tuition & budgets','Invoices; scholarships; deadlines; savings targets; funding scenarios; category budgets; credit checklist.'),
 ('transfer','Transfers & payments','FX/provider comparison; locked quotes; sends, deposits, top-ups, withdrawals; beneficiaries; receipts.'),
 ('ai','AI copilot','Chat and tool history; funding/budget advice; follow-ups; navigation; reviewed actions; model selection.'),
 ('family','Family & alerts','Invitations; recipient/transfer approvals; reminders; automatic alerts; snooze/dismiss; activity and ledger.'),
 ('document','Account & preferences','Plans/quotas; checkout/history; saved cards/models; profile/security; reports; data export and closure.')]
for i,(ic,a,b) in enumerate(student):feature(s,ic,a,b,5.35+(i%2)*3.72,2.12+(i//2)*1.52,3.52,3+i)
text(s,'English / Bengali / Chinese   ·   Themes   ·   Search   ·   Country & currency',5.83,6.76,6.8,.23,10,TEAL)
note(s,6,'The visible labels cover the student feature families; original detailed inventory is retained above. Manual external channels include bank/wire, PayPal, Alipay, WeChat Pay and UnionPay, alongside internal wallets and demo cards. The native doughnut and animated dashboard are illustrative. Profile/security includes photo/name/email, password change/reset, timezone, notification settings and account closure; subscriptions include monthly/annual sandbox purchase, limits, cancellation and history.')

# 7 parent
s=new('Parent features','Features / Parent')
art(s,'parent',7,.65,2.04,4.22,3.165)
text(s,'Support with guardrails.',.72,5.47,4.15,.55,24,INK,True,reveal=3)
text(s,'Platform ceiling',.74,6.19,1.5,.25,11,GREY);text(s,'Parent limit',2.49,6.19,1.5,.25,11,GREY)
shape(s,.74,6.59,1.49,.10,BLUE,MSO_SHAPE.RECTANGLE);shape(s,2.49,6.59,.98,.10,TEAL,MSO_SHAPE.RECTANGLE)
parent=[
 ('family','Family circle','Invite/link students; manage or revoke relationships; review linked-student finances.'),
 ('globe','Visibility controls','Balances; categories; descriptions; tuition; budgets; permission-aware alerts.'),
 ('check','Approval center','Approve/reject requests and beneficiaries; reasons/history; tuition and conversion sign-off.'),
 ('shield','Spending limits','Thresholds; per-transfer/daily/monthly caps; categories; auto-approval; one-time recipients.'),
 ('wallet','Parent money & AI','Own wallets; fund students; deposits/payments; transactions/CSV; beneficiaries; reviewed AI actions.'),
 ('document','Shared account tools','Plans; saved cards/models; reports; alerts; profile/security; preferences; export and closure.')]
for i,(ic,a,b) in enumerate(parent):feature(s,ic,a,b,5.35+(i%2)*3.72,2.12+(i//2)*1.52,3.52,3+i)
text(s,'The strictest applicable limit wins.',5.83,6.76,6.8,.23,10,TEAL)
note(s,7,'Illustrated sliders are policy controls, not measured values. Parents can only tighten platform ceilings. Multiple parent policies use the strictest applicable cap. Shared account features include localization/themes, preferences, subscription history/cancellation and data export. Parent AI cannot approve requests by itself.')

# 8 admin
s=new('Admin features','Features / Admin & moderator')
art(s,'admin',8,.65,2.04,4.22,3.165)
text(s,'Operate with evidence.',.72,5.47,4.15,.55,24,INK,True,reveal=3)
native_chart(s,'column',.64,6.02,3.9,.85,['Queued','Processing','Done'],[4,2,8],[GOLD,BLUE,TEAL],order=4)
text(s,'Illustrative settlement queue',.75,6.83,4,.18,9,GREY)
admin=[
 ('family','Users & moderation','Users/roles; suspend/reactivate; metrics; delegated moderator sections; account oversight.'),
 ('transfer','Settlements','Review/process/complete/fail; payout refs; held funds/refunds; channels and receiving details.'),
 ('shield','Financial policy','Risk/frequency/review limits; FX rates; providers, fees, spreads and routes; credit guidance.'),
 ('globe','Catalog & plans','Countries/currencies; defaults/settings; subscription pricing, quotas and feature limits.'),
 ('ai','AI & team','Shared models/keys; connection tests; run/tool logs; published team profiles and permissions.'),
 ('document','Reports & evidence','Resolve reports; audit logs; transfer records/receipts; project and security evidence exports.')]
for i,(ic,a,b) in enumerate(admin):feature(s,ic,a,b,5.35+(i%2)*3.72,2.12+(i//2)*1.52,3.52,3+i)
text(s,'Moderators receive access only to assigned sections.',5.83,6.76,6.8,.23,10,TEAL)
note(s,8,'The three queue totals are illustrative and do not come from a running database. Team features include publishing, designations, CEO selection and limited moderator self-editing. Moderation scopes: reports, team, agent, models, users, finance and catalog.')

# 9
s=new('Security & privacy, at every step.','Security & Privacy')
art(s,'security',9,.65,2.04,6.0,4.5)
security=[('Identity','Hashed passwords; signed sessions\nand revocation.'),('Access','Role, ownership and family visibility\nchecked on the server.'),('Data & AI','Encrypted model keys; validated tools;\ncredentials entered in secure forms.'),('Money integrity','Confirmation, expiring tokens, locks,\nidempotency and balanced ledger.')]
for i,(a,b) in enumerate(security):point(s,i+1,a,b,7.02,2.09+i*1.08,5.56,3+i)
text(s,'Visibility settings  ·  Data export  ·  Account closure',.73,6.68,8.3,.27,12,TEAL,True)
text(s,'Sandbox payments',9.71,6.68,2.9,.27,12,GREY,False,PP_ALIGN.RIGHT)
note(s,9,'The shield diagram shows layered code-level controls. This is not independent certification. AI providers may receive selected financial/tool context. No full PAN/CVV is stored. Profile photos are validated. Account closure is not an erasure promise. No assertion of encryption of all financial data at rest is made.')

# 10
s=new('The people behind FinBridge AI','Developer introduction')
# editable normal pictures, with circular crop preserved
for i in range(2):
 x=1.32+i*6.57
 p=picture(s,BASE/f'demo-developer-{i+1}.png',x+.64,2.17,2.36,2.36,2+i)
 text(s,f'Developer {i+1}',x,4.88,3.7,.55,29,INK,True,PP_ALIGN.CENTER,reveal=4+i)
 text(s,'AI & Backend' if i==0 else 'Frontend & UX',x,5.63,3.7,.45,18,TEAL,False,PP_ALIGN.CENTER,reveal=4+i)
art(s,'team',10,5.46,3.02,2.38,1.79,6)
text(s,'Demo portraits, names and roles — replace with your team.',.70,6.64,11.9,.30,12,GREY,False,PP_ALIGN.CENTER)
note(s,10,'User explicitly requested demo details to replace manually. Suggested roles are placeholders, not verified contributor roles. Replace portraits with Change Picture; crop any replacement to an oval. Names and roles remain editable text.')

# Simple slide transitions; PowerPoint automation adds validated native effects.
for s in P.slides:
 tr=OxmlElement('p:transition');tr.set('spd','med');tr.set('advClick','1');tr.append(OxmlElement('p:fade'));s._element.append(tr)
P.save('presentation-deck.pptx')
print('Created visual redesign:',len(P.slides),'slides,',Path('presentation-deck.pptx').stat().st_size,'bytes')
